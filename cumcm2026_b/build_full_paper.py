"""Assemble the manuscript from verified evaluation reports and editable equations."""
from __future__ import annotations

import json
import re
import statistics
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
WORK = ROOT / 'output' / 'full_paper'
sys.path.insert(0, str(ROOT.parent / 'tmp' / 'paper_deps'))
from latex2mathml.converter import convert
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCES: list[dict] = []


def load(relative):
    path = ROOT / relative
    obj = json.loads(path.read_text(encoding='utf-8'))
    SOURCES.append({'path': relative, 'fields': list(obj)})
    return obj


def table(headers, rows):
    def cell(value):
        return str(value).replace('\n','<br>')
    return '\n'.join(['| ' + ' | '.join(map(cell,headers)) + ' |',
                      '| ' + ' | '.join(['---'] * len(headers)) + ' |',
                      *['| ' + ' | '.join(map(cell, row)) + ' |' for row in rows]])


def fmt(x, n=2):
    return f'{x:.{n}f}'


def pct(x, n=2):
    return f'{100*x:.{n}f}%'


def count_tables(prefix):
    records = []
    groups = []
    for n in range(10, 17):
        d = load(f'artifacts/benchmarks/{prefix}_count{n}_100.json')
        assert d['episodes'] == 100 and d['fixed_source_count'] == n
        groups.append(d)
        records.append([n, pct(d['source_clear_rate'], 4), pct(d['complete_case_rate']),
                        fmt(d['mean_episode_mean_time_s']),
                        fmt(d['mean_virtual_time_s_per_true_source']*n),
                        fmt(d['mean_travel_m']/1000)])
    c = sum(round(d['source_clear_rate']*d['episodes']*d['fixed_source_count']) for d in groups)
    n = sum(d['episodes']*d['fixed_source_count'] for d in groups)
    t = sum(d['mean_virtual_time_s_per_true_source']*d['episodes']*d['fixed_source_count'] for d in groups)
    summary = (f'700局共包含{n}个源，清除{c}个；源级清除率为{pct(c/n,4)}，'
               f'整局全清率为{pct(statistics.fmean(d["complete_case_rate"] for d in groups))}。'
               f'逐局平均定位清除时间为{fmt(statistics.fmean(d["mean_episode_mean_time_s"] for d in groups))} s/已清源，'
               f'按真实源加权的时间为{fmt(t/n)} s/真实源。')
    return table(['真实源数','源级清除率','全清率','平均T/C\ns/源','整局时间\ns','路程\nkm'], records), summary


def assemble():
    metrics = load('artifacts/q12_paper/metrics.json')
    q1 = metrics['q1_example']
    replacements = {
        'Q1_TABLE': table(['评价量','数值'], [
            ['定位区域面积/m²',fmt(q1['area_m2'],3)],
            ['定位区域直径/m',fmt(q1['diameter_m'],3)],
            ['直径的一半/m',fmt(q1['diameter_m']/2,3)],
            ['最小包围圆半径/m',fmt(q1['minimum_enclosing_circle_radius_m'],3)],
            ['直径圆能否覆盖','不能']]),
        'Q2_TABLE': table(['排名','第二测点/m','枚举最大直径/m','枚举平均直径/m'],
                         [[r['rank'], f"({r['x_m']:.0f},{r['y_m']:.0f})",fmt(r['worst_diameter_m'],3),fmt(r['mean_diameter_m'],3)]
                          for r in metrics['q2_example']['top10'][:5]])
    }
    paired = load('artifacts/benchmarks/q3_baseline_vs_joint_100.json')
    replacements['Q3_PAIRED_TABLE'] = table(['策略','全清率','平均T/C\ns/源','平均路程\nkm','检测次数'],
        [[name,pct(paired[key]['complete_case_rate']),fmt(paired[key]['mean_time_s_per_source']),
          fmt(paired[key]['mean_travel_m']/1000),fmt(paired[key]['mean_measures'])]
         for key,name in [('baseline','分阶段基线'),('joint','联合策略joint')]])
    for q, prefix in [(3,'q3_joint'),(4,'q4_source98')]:
        replacements[f'Q{q}_COUNT_TABLE'], replacements[f'Q{q}_COUNT_SUMMARY'] = count_tables(prefix)

    official3 = []
    for i in range(4,11):
        d = load(f'artifacts/official_practice/q3_run_{i:03d}_joint/summary.json')
        assert d['mode'] == 'official_practice'
        official3.append(d)
    replacements['Q3_OFFICIAL_TABLE'] = table(['演练案例编码','清除数/真实数','平均T/C\ns/源','官方运行时间\ns'],
        [[d['official_statistics']['case_code'],f"{d['cleared_count']}/{d['official_statistics']['jammer_count']}",
          fmt(d['virtual_time_s']/d['cleared_count']),fmt(d['official_statistics']['program_run_duration_ms']/1000,3)]
         for d in official3])
    total_t = sum(d['virtual_time_s'] for d in official3)
    total_c = sum(d['cleared_count'] for d in official3)
    replacements['Q3_OFFICIAL_SUMMARY'] = (
        f'七局均全清，累计清除{total_c}/{sum(d["official_statistics"]["jammer_count"] for d in official3)}个源。'
        f'逐局T/C均值为{fmt(statistics.fmean(d["virtual_time_s"]/d["cleared_count"] for d in official3))} s/源，'
        f'累计总时间除以累计源数为{fmt(total_t/total_c)} s/源。'
        '最快样本run010为15/15全清，虚拟总时间3435.957 s、平均229.06 s/源，官方程序运行时间2.418 s。'
        '这一局是最佳演练案例展示，七局整体结果用于说明重复运行表现。')

    oldrows=[]
    for dist,label in [('fixed16','固定16源'),('mixed','混合源')]:
        for key, name in [('base','原基线'),('geometry','conservative'),('balanced','balanced')]:
            d=load(f'artifacts/benchmarks/q4_final_{key}_{dist}_500.json')
            oldrows.append([label,name,pct(d['source_clear_rate'],4),pct(d['complete_case_rate']),
                            fmt(d['mean_episode_mean_time_s'])])
    replacements['Q4_PAIRED_TABLE']=table(['场景','配置','源级清除率','全清率','平均T/C\ns/源'],oldrows)
    riskrows=[]
    for file,dist,name in [
        ('q4_chance_validation_mix010_300.json','混合源','v2 α=0.10'),
        ('q4_chance_validation_fix010_300.json','固定16源','v2 α=0.10'),
        ('q4_v2_alpha025_mix_300.json','混合源','source-98'),
        ('q4_v2_alpha025_fix_300.json','固定16源','source-98')]:
        d=load('artifacts/benchmarks/'+file)
        riskrows.append([dist,name,pct(d['source_clear_rate'],4),pct(d['complete_case_rate']),
                         fmt(d['mean_episode_mean_time_s']),fmt(d['mean_virtual_time_s_per_true_source'])])
    replacements['Q4_RISK_TABLE']=table(['场景','配置','源级清除率','全清率','平均T/C\ns/已清源','ΣT/ΣN\ns/真实源'],riskrows)
    off4=[]
    for i,suffix,name in [(9,'source_efficient','v1'),(10,'source_efficient','v1'),
                          (11,'source_efficient','v1'),(12,'source_efficient','v1'),
                          (13,'source_efficient_v2','v2'),(14,'source_efficient_v2','v2'),
                          (15,'source_98','source-98')]:
        d=load(f'artifacts/official_practice/q4_run_{i:03d}_{suffix}/summary.json')
        assert d['mode']=='official_practice'
        s=d['official_statistics']
        off4.append([s['case_code'],name,f"{d['cleared_count']}/{s['jammer_count']}",
                     fmt(d['virtual_time_s']/d['cleared_count']),fmt(s['program_run_duration_ms']/1000,3)])
    replacements['Q4_OFFICIAL_TABLE']=table(['演练案例编码','配置','清除/真实','平均T/C\ns/源','官方运行\ns'],off4)
    replacements['EVIDENCE_TABLE']=table(['结果','支撑材料相对路径'],[
        ['Q1与Q2构造例','artifacts/q12_paper/metrics.json'],
        ['Q3同种子比较','artifacts/benchmarks/q3_baseline_vs_joint_100.json'],
        ['Q3源数分层','artifacts/benchmarks/q3_joint_count{10…16}_100.json'],
        ['Q4旧配对实验','artifacts/benchmarks/q4_final_{base,geometry,balanced}_{fixed16,mixed}_500.json'],
        ['Q4双风险v2','artifacts/benchmarks/q4_chance_validation_{mix,fix}010_300.json'],
        ['Q4速度配置','artifacts/benchmarks/q4_v2_alpha025_{mix,fix}_300.json'],
        ['Q4源数分层','artifacts/benchmarks/q4_source98_count{10…16}_100.json'],
        ['官方演练','artifacts/official_practice/下对应run目录的summary.json及客户端日志'],
        ['神经停止消融','artifacts/neural_belief_v1/test_cpu.json及history.json']])
    text=(WORK/'manuscript_source.md').read_text(encoding='utf-8')
    for key,value in replacements.items():
        text=text.replace('{{'+key+'}}',value.replace('\nkm','<br>km').replace('\ns/','<br>s/').replace('\ns|','<br>s|').replace('\ns ','<br>s '))
    # Keep table rows single-line; line breaks are encoded as HTML only inside cells.
    text=text.replace('时间\ns','时间<br>s').replace('运行\ns','运行<br>s')
    assert not re.search(r'\{\{[A-Z_0-9]+\}\}',text)
    captions=['主要符号与单位','问题一构造算例结果','问题二排名前五的候选点',
              '问题三同种子策略比较','问题三joint源数分层结果','问题三官方演练结果',
              '问题3正式测试结果','问题四同批候选配置比较','问题四双风险配置独立验证',
              '问题四相关官方演练结果','问题4正式测试结果','问题四source-98源数分层结果',
              '代码入口与实现对应','结果与数据文件对应']
    numbered=[];count=0;in_table=False
    for line in text.splitlines():
        if line.startswith('表 问题'):
            continue
        if line.startswith('|') and not in_table:
            count+=1;numbered.extend([f'表{count} {captions[count-1]}',''])
        in_table=line.startswith('|')
        numbered.append(line)
    assert count==len(captions)
    text='\n'.join(numbered)+'\n'
    (WORK/'2026B题_完整论文初稿.md').write_text(text,encoding='utf-8')
    (WORK/'data_provenance.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2),encoding='utf-8')
    return text


TRANSFORM=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))


def math_element(latex):
    mathml=convert(latex)
    element=TRANSFORM(etree.fromstring(mathml.encode('utf-8'))).getroot()
    for r in element.findall('.//'+qn('m:r')):
        rpr=OxmlElement('w:rPr')
        fonts=OxmlElement('w:rFonts')
        for attr in ('ascii','hAnsi','eastAsia'):
            fonts.set(qn('w:'+attr),'Cambria Math')
        rpr.append(fonts)
        r.append(rpr)
    return element


def inline(paragraph,text):
    text=text.replace('<br>','\n')
    for part in re.split(r'(\$[^$]+\$)',text):
        if part.startswith('$') and part.endswith('$'):
            paragraph._p.append(math_element(part[1:-1]))
        elif part:
            paragraph.add_run(part.replace('`',''))


def set_style(style,size,east='宋体',bold=False):
    style.font.name='Times New Roman'
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),east)


def make_docx(text):
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.35);sec.bottom_margin=Cm(2.15)
    sec.left_margin=Cm(2.4);sec.right_margin=Cm(2.4)
    sec.header_distance=Cm(1);sec.footer_distance=Cm(1)
    set_style(doc.styles['Normal'],11)
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=1.28;normal.space_after=Pt(5)
    normal.first_line_indent=Pt(22);normal.widow_control=True
    set_style(doc.styles['Title'],18,'黑体',True)
    for n,size in [(1,15),(2,12),(3,11)]:
        style=doc.styles[f'Heading {n}']
        set_style(style,size,'黑体',True)
        style.paragraph_format.space_before=Pt(12)
        style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.first_line_indent=Pt(0)
        style.paragraph_format.keep_with_next=True
    f=sec.footer.paragraphs[0]
    f.paragraph_format.first_line_indent=Pt(0);f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    lines=text.splitlines();i=0;code=False;eq_count=0;tables=0
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('```'):
            code=not code;i+=1;continue
        if code:
            p=doc.add_paragraph(line)
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.left_indent=Cm(0.3)
            p.paragraph_format.space_after=Pt(2)
            for r in p.runs:r.font.size=Pt(10)
        elif line.startswith('# '):
            p=doc.add_paragraph(line[2:],style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Pt(0)
        elif line.startswith('## '):
            title=line[3:];p=doc.add_paragraph(title,style='Heading 1')
            if title=='1 问题重述':p.paragraph_format.page_break_before=True
            if title=='摘要':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('$$'):
            formula=line.strip('$').strip()
            tag=re.search(r'\\tag\{([^}]+)\}',formula)
            formula=re.sub(r'\\tag\{[^}]+\}','',formula).strip()
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.keep_together=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p._p.append(math_element(formula))
            if tag:p.add_run('  ('+tag[1]+')')
            eq_count+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.fullmatch(r'[| :\-]+',lines[i]):
                    rows.append([cell.strip() for cell in lines[i].strip().strip('|').split('|')])
                i+=1
            cols=len(rows[0]);assert all(len(r)==cols for r in rows),(rows,cols)
            t=doc.add_table(rows=0,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            if '编码' in rows[0][0]:widths=[5.2,3.2,3.6,4.2] if cols==4 else [4.6,2.2,2.2,3.7,3.5]
            elif cols==2:widths=[4.2,12.0]
            elif cols==3:widths=[3.1,10.6,2.5] if rows[0][0]=='符号' else [3.0,5.3,7.9]
            elif cols==6:widths=[2.2,2.8,2.7,2.5,3.1,2.9]
            else:widths=[16.2/cols]*cols
            for col,w in zip(t.columns,widths):col.width=Cm(w)
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            t._tbl.tblPr.append(borders)
            for ri,row in enumerate(rows):
                cells=t.add_row().cells
                trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if ri==0:trpr.append(OxmlElement('w:tblHeader'))
                for ci,(cell,value,w) in enumerate(zip(cells,row,widths)):
                    cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    if len(rows)<=9 and ri<len(rows)-1:p.paragraph_format.keep_with_next=True
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if cols==2 or (cols==3 and ci==1) else WD_ALIGN_PARAGRAPH.CENTER
                    inline(p,value)
                    for r in p.runs:
                        r.font.size=Pt(9.5 if cols>=5 else 10)
                        r.bold=ri==0
                    if ri==0:
                        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EFEFEF');cell._tc.get_or_add_tcPr().append(shd)
            tables+=1
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=0.3
            continue
        else:
            p=doc.add_paragraph();inline(p,line)
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            following=next((s for s in lines[i+1:] if s.strip()),'')
            if following.startswith('$$'):p.paragraph_format.keep_with_next=True
            if line.startswith('【图'):
                p.paragraph_format.first_line_indent=Pt(0)
                for r in p.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('555555')
            elif re.match(r'^表\d+ ',line):
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            elif re.match(r'^\[\d+\]',line):
                p.paragraph_format.first_line_indent=Pt(0)
                for r in p.runs:r.font.size=Pt(9.5)
        i+=1
    out=WORK/'2026B题_完整论文初稿.docx'
    for root in (doc.styles.element,doc.element):
        for border in root.findall('.//'+qn('w:pBdr')):
            border.getparent().remove(border)
    doc.save(out)
    print(json.dumps({'docx':str(out),'characters':len(text),'display_equations':eq_count,'tables':tables,'paragraphs':len(doc.paragraphs)},ensure_ascii=False))


if __name__=='__main__':
    make_docx(assemble())
