from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REF = Path(r"C:\Users\10399\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-design-report\assets\reference.docx")
ART = ROOT / "artifacts" / "q12_paper"
METRICS = json.loads((ART / "metrics.json").read_text(encoding="utf-8"))
OUT_DIR = ROOT / "output"
OUT = OUT_DIR / "2026B题_问题一与问题二_建模解答.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, latin="Helvetica Neue", east="宋体", size=None, bold=None, color=None) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, latin, east, size, bold=False, color="000000") -> None:
    style.font.name = latin
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color="666666")


def add_bottom_border(paragraph, color="BFBFBF", size="6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_keep_lines(paragraph) -> None:
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True


def add_body(doc, text: str, *, first_indent=True, bold_lead: str | None = None, citation=None):
    p = doc.add_paragraph(style="normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(4)
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, east="黑体", bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if citation:
        r = p.add_run(citation)
        set_run_font(r, size=9)
    set_keep_lines(p)
    return p


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style="normal")
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.6)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run("•  " + text)
    set_run_font(r)
    set_keep_lines(p)
    return p


def add_numbered(doc, n: str, title: str, text: str):
    p = doc.add_paragraph(style="normal")
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(f"{n}  {title}。")
    set_run_font(r, east="黑体", bold=True)
    r2 = p.add_run(text)
    set_run_font(r2)
    set_keep_lines(p)
    return p


def add_heading(doc, text: str, level=1, page_break=False):
    style_name = f"Heading {level}"
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    p.paragraph_format.space_before = Pt(11 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    r = p.add_run(text)
    set_run_font(r, east="黑体", size=22 if level == 1 else (14 if level == 2 else 11), bold=level >= 2)
    return p


def add_omath(paragraph, text: str) -> None:
    math = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mrpr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    mrpr.append(sty)
    mt = OxmlElement("m:t")
    mt.text = text
    mr.extend([mrpr, mt])
    math.append(mr)
    paragraph._p.append(math)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def set_table_borders(table, color="B7B7B7", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def add_equation(doc, text: str, number: int):
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Cm(0.6), Cm(13.1), Cm(1.5)]
    for c, w in zip(t.rows[0].cells, widths):
        c.width = w
        set_cell_margins(c, 20, 10, 20, 10)
    remove_table_borders(t)
    p = t.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_omath(p, text)
    n = t.cell(0, 2).paragraphs[0]
    n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = n.add_run(f"({number})")
    set_run_font(r, east="宋体", size=9)
    prevent_row_split(t.rows[0])
    return t


def add_caption(doc, text: str, note: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = bool(note)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run_font(r, east="黑体", size=9.5, bold=True)
    if note:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(6)
        rr = p2.add_run(note)
        set_run_font(rr, east="宋体", size=8.5, color="666666")
        set_keep_lines(p2)
    return p


def add_figure(doc, image_path: Path, width_cm: float, caption: str, note: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    add_caption(doc, caption, note)


def add_table(doc, headers, rows, widths=None, font_size=9.2, caption=None):
    if caption:
        add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        if widths:
            cell.width = Cm(widths[j])
        set_cell_shading(cell, "E7E7E7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        set_run_font(r, east="黑体", size=font_size, bold=True)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            if widths:
                cells[j].width = Cm(widths[j])
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 25 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
        prevent_row_split(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    set_style_font(styles["normal"], "Helvetica Neue", "宋体", 10.5, False, "111111")
    normal_pf = styles["normal"].paragraph_format
    normal_pf.space_after = Pt(4)
    normal_pf.line_spacing = 1.45
    set_style_font(styles["Title"], "Helvetica Neue", "黑体", 34, True, "000000")
    set_style_font(styles["Subtitle"], "Helvetica Neue", "微软雅黑", 13, False, "595959")
    for name, latin, east, size, bold in [
        ("Heading 1", "Helvetica Neue", "黑体", 22, False),
        ("Heading 2", "Helvetica Neue", "黑体", 14, True),
        ("Heading 3", "Helvetica Neue", "黑体", 11, True),
    ]:
        set_style_font(styles[name], latin, east, size, bold, "000000")
    if "Small note" not in [s.name for s in styles]:
        s = styles.add_style("Small note", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(s, "Helvetica Neue", "宋体", 8.5, False, "666666")


def configure_section(section, *, cover=False) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    if cover:
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.header_distance = Cm(0.5)
        section.footer_distance = Cm(0.5)
    else:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(REF))
    clear_document_body(doc)
    configure_styles(doc)

    cover_sec = doc.sections[0]
    configure_section(cover_sec, cover=True)
    cover_sec.different_first_page_header_footer = True

    # Cover: source-derived hierarchy with a problem-specific figure.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("2026 CUMCM · PROBLEM B")
    set_run_font(r, size=9, bold=True, color="666666")
    add_bottom_border(p, color="8C8C8C", size="8")

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("有界示向误差下的无线电干扰源\n交会定位与稳健观测点设计")
    set_run_font(r, east="黑体", size=31, bold=True)
    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("2026年高教社杯全国大学生数学建模竞赛 B题 · 问题一与问题二")
    set_run_font(r, east="微软雅黑", size=12.5, color="595959")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(ART / "q1_localization.png"), width=Cm(17.7))

    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    remove_table_borders(t)
    vals = [
        ("研究对象", "全向干扰源"),
        ("核心方法", "凸几何与稳健优化"),
        ("文档状态", "问题一、二完整解答"),
    ]
    for cell, (label, value) in zip(t.rows[0].cells, vals):
        cell.width = Cm(5.7)
        set_cell_margins(cell, 80, 100, 80, 100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + "\n")
        set_run_font(r, east="微软雅黑", size=8, bold=True, color="777777")
        r = p.add_run(value)
        set_run_font(r, east="黑体", size=10.5, bold=True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("阶段论文稿  ·  2026年9月")
    set_run_font(r, east="微软雅黑", size=9.5, color="666666")

    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(sec, cover=False)
    sec.different_first_page_header_footer = False
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), "1")
    sec._sectPr.append(pg_num)
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(3)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(14.0))
    r = hp.add_run("无线电干扰源交会定位与稳健观测点设计")
    set_run_font(r, east="微软雅黑", size=8.2, color="666666")
    r = hp.add_run("\t2026 B题")
    set_run_font(r, east="微软雅黑", size=8.2, color="666666")
    add_bottom_border(hp, color="BFBFBF", size="4")
    add_page_field(sec.footer.paragraphs[0])

    # Abstract.
    add_heading(doc, "摘要", 1)
    add_body(doc, "本文针对无线电干扰源示向读数具有±1°有界误差、重复测量不能消除定点偏差且干扰源最大有效接收半径未知的条件，建立基于凸可行域的确定性定位模型，并完成问题一、问题二的算法设计与数值验证。模型不把误差擅自概率化，而把每次读数直接转化为一个方向扇形约束；由目标圆域、各扇形与最大接收圆盘的交集得到所有与观测一致的位置集合。")
    add_body(doc, "针对问题一，将圆弧边界作可控精度的正多边形近似，再采用半平面裁剪求交会定位多边形。由于所得集合为凸集，其直径必由两个顶点取得，故使用旋转卡壳在线性时间内搜索对踵点。进一步通过最小包围圆判定“直径为定位区域直径的圆”能否覆盖该区域：当且仅当最小包围圆半径不超过区域直径的一半时可以覆盖。等边三角形给出严格反例，因此一般情况下答案是否定的。构造算例得到定位区域直径40.748 m、最小包围圆半径22.600 m，而直径圆半径仅20.374 m，验证了判据。")
    add_body(doc, "针对问题二，在首次观测可行域内离散干扰源假设，并对候选第二测点枚举全部有界示向误差。以“漏测比例—最坏后验直径—平均后验直径—行驶距离”为字典序目标，先保证在1000 m最小有效半径下尽可能接收到信号，再压缩最坏定位误差，最后兼顾效率。该策略把经典的近正交交会几何推广为含覆盖不确定性的稳健选址。标准化算例中，1009个候选点有51个实现零漏测，推荐区域由其中最坏直径处于前20%的11个点构成；最优点为(1000,100) m，最坏后验直径110.854 m。交会角从10°增至90°时，后验直径由354.620 m降至42.626 m，说明候选区域应优先布置在首条示向线两侧、可形成60°—90°锐交会角且满足可靠接收的带状区域。")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run("关键词：")
    set_run_font(r, east="黑体", bold=True)
    r = p.add_run("示向定位；有界误差；凸可行域；旋转卡壳；最小包围圆；稳健选址")
    set_run_font(r)

    add_heading(doc, "1  问题重述", 1, page_break=True)
    add_heading(doc, "1.1  问题背景", 2)
    add_body(doc, "在半径1800 m的圆形目标区域内，移动机器人需要利用无线电示向测量定位并清除干扰源。全向干扰源的最大有效接收半径位于1000—1500 m之间；机器人处于有效范围且距源超过5 m时，可读出由正东方向起、逆时针为正的方位角。题面规定，同一位置上的读数误差固定在[−1°,1°]内，因此在原地重复测量不能带来独立样本，也不能通过平均缩小误差。")
    add_heading(doc, "1.2  两个任务", 2)
    add_numbered(doc, "（1）", "定位区域直径", "由若干检测点坐标及其示向角，构造干扰源的交会定位多边形，计算区域内任意两点的最大距离；并判断直径等于该最大距离的圆能否覆盖整个定位区域。")
    add_numbered(doc, "（2）", "第二测点选择", "在已知一个检测点及其示向角后，给出能获得较好定位结果的第二检测点选择策略，并给出可供机器人执行的候选区域。")

    add_heading(doc, "2  问题分析", 1)
    add_heading(doc, "2.1  关键困难", 2)
    add_bullet(doc, "误差是确定性的区间误差而非随机白噪声，单点观测对应的是无限延伸的角扇形，而不是一条射线。")
    add_bullet(doc, "目标区域与接收半径约束含圆弧；若直接对任意曲边区域穷举点对，既慢又难以保证直径精度。")
    add_bullet(doc, "第二测点不仅要形成较大交会角，还要在干扰源实际半径可能只有1000 m时仍能接收到信号，因此“几何精度最好”与“保证检测成功”可能冲突。")
    add_bullet(doc, "题目要求候选区域而非单个理想点，故选址结果应包含可实施的区域定义、排序规则及降级方案。")
    add_heading(doc, "2.2  总体思路", 2)
    add_body(doc, "第一步把每条示向读数变为两个有向半平面的交；第二步与目标圆域及检测圆盘求交，得到凸可行域；第三步用旋转卡壳计算直径，用随机增量法计算最小包围圆。完成首次观测后，在可行域内布置干扰源假设、在目标区域内生成第二测点网格，并在最不利误差与最小可靠接收半径下逐点评价。该流程全部使用可解释的几何量，且可以直接嵌入后续机器人路径规划。")

    add_heading(doc, "3  模型假设", 1)
    assumptions = [
        "机器人位置坐标已校准，位置误差相对示向误差可忽略；若后续获得定位误差上界，可按同样方法把检测点扩张为小圆盘。",
        "各次示向读数均满足题面给定的±1°硬上界，不假设误差服从正态分布，也不把同点重复观测视为独立样本。",
        "问题一、二只研究全向干扰源；只要检测点到真实源的距离处于(5 m,R]且频道正确，就能得到方位角，其中R∈[1000,1500] m。",
        "对已经得到示向角的检测点，真实源必在该点1500 m以内；为保证第二测点对所有可能R都有效，设计阶段采用1000 m作为可靠接收半径。",
        "圆弧使用足够密的内接正多边形离散，离散误差由弦高控制并在敏感性分析中检查。",
    ]
    for i, text in enumerate(assumptions, 1):
        add_numbered(doc, str(i), "假设", text)

    add_heading(doc, "4  符号说明", 1)
    symbols = [
        ("O", "目标圆域中心", "m"),
        ("R₀", "目标区域半径，R₀=1800", "m"),
        ("Sᵢ", "第i个检测点坐标", "m"),
        ("θ̂ᵢ", "第i个检测点的示向读数", "°或rad"),
        ("ε", "示向误差上界，ε=1°", "°或rad"),
        ("Wᵢ", "第i次观测对应的方向扇形", "—"),
        ("Ω", "所有观测共同确定的定位可行域", "m²区域"),
        ("D(Ω)", "定位区域直径", "m"),
        ("R*", "定位区域最小包围圆半径", "m"),
        ("α", "两条视线的锐交会角", "°或rad"),
        ("pₘᵢₛₛ", "第二测点在假设集上的漏测比例", "—"),
    ]
    add_table(doc, ["符号", "含义", "单位"], symbols, widths=[2.6, 10.3, 2.2], caption="表1  主要符号")

    add_heading(doc, "5  问题一  交会定位区域及其直径", 1, page_break=True)
    add_heading(doc, "5.1  单次示向的扇形约束", 2)
    add_body(doc, "记第i个检测点为Sᵢ=(xᵢ,yᵢ)，读数为θ̂ᵢ。误差上界ε=1°，则真实方向只能落在[θ̂ᵢ−ε,θ̂ᵢ+ε]内。定义左右边界方向向量dᵢ⁻、dᵢ⁺，用二维叉积判断未知点X位于两条边界射线之间。")
    add_equation(doc, "dᵢ⁻=(cos(θ̂ᵢ−ε), sin(θ̂ᵢ−ε)),   dᵢ⁺=(cos(θ̂ᵢ+ε), sin(θ̂ᵢ+ε))", 1)
    add_equation(doc, "Wᵢ={X: cross(dᵢ⁻,X−Sᵢ)≥0,  cross(dᵢ⁺,X−Sᵢ)≤0}", 2)
    add_body(doc, "以上方向约定在扇形不跨越±π时直接成立；程序实现先把角度归一化，再统一转换为有向半平面，因而能够处理0°附近的跨界情况。由于已经获得示向角，真实源还满足‖X−Sᵢ‖≤1500 m；距源不超过5 m时题面规定无法示向而可直接清除，因此本问的有效观测天然排除了该退化情形。")

    add_heading(doc, "5.2  多观测可行域", 2)
    add_body(doc, "目标区域Ω₀是以原点为中心、半径1800 m的圆盘。综合m次观测，交会定位区域定义为目标圆域、全部示向扇形和最大接收圆盘的交。")
    add_equation(doc, "Ω=Ω₀ ∩ (⋂ᵢ₌₁ᵐ Wᵢ) ∩ (⋂ᵢ₌₁ᵐ B(Sᵢ,1500))", 3)
    add_body(doc, "圆盘和半平面均为凸集，故Ω仍为凸集。计算时把目标圆和接收圆用正N边形近似，并依次执行Sutherland—Hodgman半平面裁剪。若半径为r的圆用正N边形逼近，最大弦高为r[1−cos(π/N)]；据此可反推N，使几何离散误差低于预设阈值。若交集为空，应首先检查角度单位、频道对应和观测一致性，而不能输出虚假位置。")
    add_equation(doc, "eₙ=r[1−cos(π/N)]", 4)

    add_heading(doc, "5.3  定位区域直径算法", 2, page_break=True)
    add_body(doc, "设裁剪得到按逆时针排列的凸多边形顶点v₀,…,vₙ₋₁。凸集的直径一定由极点取得；对多边形而言，极点是顶点，因此无需在区域内部搜索。先固定一条边(vᵢ,vᵢ₊₁)，移动对踵点指针j，直到三角形面积不再增大；每个指针至多绕多边形一周，即可在线性时间枚举所有候选点对。该方法即旋转卡壳[5]。")
    add_equation(doc, "D(Ω)=max₀≤p<q<n ‖vₚ−v_q‖₂", 5)
    add_table(doc, ["步骤", "操作", "输出或判据"], [
        ("1", "构造Ω₀及各接收圆的正多边形", "弦高≤给定几何容差"),
        ("2", "按每条示向边界逐次裁剪", "凸多边形顶点序列"),
        ("3", "删除近重复点并统一为逆时针", "稳定的凸多边形"),
        ("4", "旋转卡壳枚举对踵点", "最大平方距离及端点"),
        ("5", "对平方距离开方", "定位区域直径D"),
    ], widths=[1.4, 8.1, 5.6], caption="表2  问题一直径计算流程")

    add_heading(doc, "5.4  直径圆能否覆盖定位区域", 2)
    add_body(doc, "“圆的直径等于定位区域直径”只规定圆的半径为D/2，并未保证存在合适圆心。令R*为Ω的最小包围圆半径，则半径D/2的圆可以覆盖Ω，当且仅当R*≤D/2。另一方面，任何覆盖圆都必须容纳一对直径端点，故必有R*≥D/2；因此精确条件可写成R*=D/2。最小包围圆由两个对径支撑点或三个边界点唯一确定，可用Welzl随机增量算法在线性期望时间内求得[6]。")
    add_equation(doc, "R*=min_c max_{X∈Ω} ‖X−c‖₂", 6)
    add_equation(doc, "存在直径为D的覆盖圆  ⇔  R*≤D/2  ⇔  R*=D/2", 7)
    add_body(doc, "一般情况下答案是否定的。取边长为D的等边三角形作为定位多边形，其区域直径就是D，但外接圆亦为最小包围圆，半径为D/√3>D/2，因此任何直径为D的圆都无法同时覆盖三个顶点。若最小包围圆恰由一对相距D的顶点支撑，则圆心为该点对中点、半径为D/2，此时覆盖成立。")

    add_heading(doc, "5.5  数值验证", 2)
    q1 = METRICS["q1_example"]
    add_body(doc, "为验证算法，构造真实源(500,400) m以及三个检测点(−500,0)、(400,−600)、(−200,1000) m，在真实方位角上分别加入0.6°、−0.7°、0.3°的固定误差。该算例仅用于验证模型与代码，不是题面提供的正式测试数据。")
    add_figure(doc, ART / "q1_localization.png", 15.4, "图1  三次有界示向观测的交会定位区域", "左图给出目标域和三条示向中心线；右下角放大图比较直径端点、直径圆与最小包围圆。")
    add_table(doc, ["指标", "结果", "含义"], [
        ("可行域顶点数", f"{q1['vertices']}", "交集为凸四边形"),
        ("可行域面积", f"{q1['area_m2']:.3f} m²", "示向交会后的剩余区域"),
        ("区域直径D", f"{q1['diameter_m']:.3f} m", "旋转卡壳结果"),
        ("D/2", f"{q1['diameter_m']/2:.3f} m", "题设直径圆的半径"),
        ("最小包围圆半径R*", f"{q1['minimum_enclosing_circle_radius_m']:.3f} m", "Welzl算法结果"),
        ("覆盖判定", "不能覆盖", "R*>D/2"),
    ], widths=[4.2, 4.2, 6.7], caption="表3  问题一构造算例结果")
    add_body(doc, "结果中R*=22.600 m，而D/2=20.374 m，两者相差2.226 m。图1可见直径圆遗漏了另一个极点，最小包围圆则覆盖全部顶点，数值结果与理论判据一致。")

    add_heading(doc, "6  问题二  第二检测点的稳健选址", 1, page_break=True)
    add_heading(doc, "6.1  从正交交会到稳健选址", 2)
    add_body(doc, "若干扰源位置已知，第二条视线与第一条视线越接近正交，角度误差在交点处被放大的程度越小。经典到达角定位和Fisher信息分析均表明，合理的观测轨迹和传感器几何能够显著改善定位精度[2–4]。但本题在第一次观测后只知道一条狭长可行带，而且实际有效半径可能只有1000 m；单纯追求90°交会，可能把机器人送到无法接收信号的位置。")
    add_equation(doc, "定位尺度 ∝ 1/|sin α|", 8)
    add_body(doc, "因此，本题第二测点必须同时回答三个问题：能否可靠收到信号、最坏情况下定位区域有多大、在相同定位质量下哪一点更省路程。我们采用字典序优化，使这些优先级显式可检验，而不是用缺乏依据的权重把不同量纲相加。")

    add_heading(doc, "6.2  首次观测后的假设集", 2)
    add_body(doc, "设首次检测点为S₁、读数为θ̂₁。按问题一方法得到首次可行域Ω₁。由于仅有一条示向信息，Ω₁通常是沿θ̂₁方向延伸的扇形带。用极坐标分层或规则网格在Ω₁内生成干扰源假设G={g₁,…,g_K}；在角度边界、近端和远端加密，以免漏掉决定最坏情形的极端位置。")
    add_equation(doc, "Ω₁=Ω₀ ∩ W₁ ∩ B(S₁,1500)", 9)

    add_heading(doc, "6.3  候选点的后验评价", 2)
    add_body(doc, "对候选第二测点s和任一假设g，若5<‖g−s‖≤1000 m，则无论实际最大有效半径在[1000,1500] m的何处，都能保证获得第二个示向角。设第二次固定误差δ∈[−ε,ε]，模拟读数为arg(g−s)+δ，后验定位区域为：")
    add_equation(doc, "Ω₂(s,g,δ)=Ω₁ ∩ W(s,arg(g−s)+δ,ε)", 10)
    add_body(doc, "若距离超过1000 m则记为漏测；若不超过5 m则机器人可直接清除，该候选点在任务层面成功，但不再需要计算第二条示向。对纯定位策略，可把它单列为“直接清除”状态。定义候选点的漏测比例、最坏后验直径和平均后验直径：")
    add_equation(doc, "pₘᵢₛₛ(s)=|{g∈G: ‖g−s‖>1000}|/|G|", 11)
    add_equation(doc, "Dmax(s)=max_{g∈G,δ∈E} D(Ω₂),   Dmean(s)=mean_{g∈G,δ∈E} D(Ω₂)", 12)
    add_body(doc, "其中E至少包含−1°、0°、1°，在最终计算中可进一步加密。候选点按下式从小到大排序：")
    add_equation(doc, "min_s  (pₘᵢₛₛ(s), Dmax(s), Dmean(s), ‖s−S₁‖)  （字典序）", 13)
    add_body(doc, "该目标首先减少漏测，再抑制最坏误差；只有前面的指标相同时，平均表现和行驶距离才参与决策。这与应急清除任务的风险偏好一致，也避免“多走一点”和“漏掉干扰源”被一个任意权重错误抵消。")

    add_heading(doc, "6.4  第二测点候选区域", 2)
    add_body(doc, "为了给出区域而非孤立点，定义零漏测候选集C₀，并在其中保留最坏直径排名前ρ的点。本文取ρ=20%，用于生成清晰且有冗余的推荐区域；实际执行时可根据障碍物和路径代价调整ρ。")
    add_equation(doc, "C₀={s∈Ω₀:pₘᵢₛₛ(s)=0},   Cρ={s∈C₀:Dmax(s)≤Qρ[Dmax(C₀)]}", 14)
    add_body(doc, "从几何上看，Cρ通常落在首条示向中心线的两侧：它既接近Ω₁的中远段以覆盖所有源假设，又与第一视线形成较大的交会角。故可先用解析规则筛选——距离Ω₁代表点不超过1000 m、与首条示向线的锐交会角处于60°—90°——再在筛选区内进行精细网格或连续优化。两侧区域应同时保留，机器人最终选择可达且路程较短的一侧。")

    add_heading(doc, "6.5  选址算法", 2)
    add_table(doc, ["阶段", "计算内容", "实现要点"], [
        ("A  首次可行域", "由S₁、θ̂₁生成Ω₁", "使用与问题一相同的半平面裁剪"),
        ("B  假设生成", "在Ω₁内分层采样G", "边界、远端、近端加密"),
        ("C  候选生成", "在目标圆域生成粗网格", "先做1000 m覆盖和交会角筛选"),
        ("D  最坏评价", "遍历g和δ求Ω₂及其直径", "δ先取端点与0，必要时加密"),
        ("E  稳健排序", "按四元组字典序排序", "覆盖优先、精度次之、路程最后"),
        ("F  区域输出", "保留Cρ并连成候选斑块", "剔除障碍后选最近可达点"),
    ], widths=[2.2, 6.0, 6.9], caption="表4  第二测点稳健选址流程")
    add_body(doc, "若C₀为空，说明不存在一个点能在1000 m半径下覆盖全部首次假设。此时不应强行给出“保证测得”的结论，而应依次最小化pₘᵢₛₛ、Dmax和行驶距离，并把剩余未覆盖的Ω₁子区作为后续搜索目标。这是算法的明确降级路径。")

    add_heading(doc, "6.6  标准化算例", 2)
    q2 = METRICS["q2_example"]
    add_body(doc, "取S₁=(0,0) m、θ̂₁=35°。沿首次扇形的半径方向每50 m采样，在角宽内取5个方向，共得到150个干扰源假设；在半径1800 m目标圆内使用100 m网格生成1009个候选第二测点。该算例用于展示策略的可执行输出，正式参赛时只需替换为题目给定的首次检测点和读数。")
    add_figure(doc, ART / "q2_candidate_region.png", 15.6, "图2  第二测点稳健评价与推荐候选区域", "绿色点为零漏测且最坏后验直径位于前20%的推荐点；星号为字典序最优点。")
    add_table(doc, ["评价量", "计算结果"], [
        ("干扰源假设数", f"{q2['hypotheses']}"),
        ("候选点数", f"{q2['evaluated_candidates']}"),
        ("零漏测候选点数", f"{q2['zero_miss_candidates']}"),
        ("推荐区域候选点数", f"{q2['recommended_region_candidates']}"),
        ("推荐区最坏直径阈值", f"{q2['recommended_worst_diameter_threshold_m']:.3f} m"),
        ("字典序最优点", f"({q2['best']['point'][0]:.0f},{q2['best']['point'][1]:.0f}) m"),
        ("最优点最坏后验直径", f"{q2['best']['worst_diameter_m']:.3f} m"),
        ("最优点平均后验直径", f"{q2['best']['mean_diameter_m']:.3f} m"),
        ("最优点行驶距离", f"{q2['best']['travel_distance_m']:.3f} m"),
    ], widths=[8.0, 7.1], caption="表5  第二测点标准化算例结果")
    add_body(doc, "最优点实现零漏测，其最坏后验直径为110.854 m。推荐区域含11个离散点，提供了绕开局部障碍或缩短真实道路距离的替代空间。若机器人只能访问网格之外的位置，可对推荐点做Delaunay邻接或半个网格间距的缓冲，得到连续候选斑块，再在斑块内局部细化。")

    add_heading(doc, "6.7  可直接执行的选点规则", 2)
    add_numbered(doc, "1", "先保接收", "以1000 m而非1500 m检查第二测点对首次可行域的覆盖；能够覆盖全部假设的点优先。")
    add_numbered(doc, "2", "再取大交会角", "在可靠接收候选中，优先选择首条示向线两侧、对首次可行域主体形成60°—90°锐交会角的位置。")
    add_numbered(doc, "3", "按最坏直径排序", "对±1°误差端点和0°逐一模拟，用问题一直径算法计算后验区域，选择最大值最小者。")
    add_numbered(doc, "4", "保留区域冗余", "输出前20%稳健点组成的候选区域；遇到障碍时在区域内选最近可达点，无需重新定义定位目标。")
    add_numbered(doc, "5", "无法全覆盖时降级", "先使漏测比例最小，并记录未覆盖子区；后续机器人搜索必须优先补扫该子区。")

    add_heading(doc, "7  敏感性与稳健性分析", 1, page_break=True)
    add_heading(doc, "7.1  交会角敏感性", 2)
    add_body(doc, "固定名义源位置与第二测点到源的距离，仅改变两条视线的锐交会角，并对两次±1°有界误差求最坏后验区域。结果显示，交会角由10°增至90°时，区域直径持续下降，且小角度区间最敏感。")
    add_figure(doc, ART / "q2_angle_sensitivity.png", 14.8, "图3  后验定位直径对交会角的敏感性", "名义源为(1000,0) m，第二测点距源700 m；曲线用于比较几何趋势。")
    sensitivity = METRICS["q2_crossing_angle_sensitivity"]
    chosen = [sensitivity[i] for i in [0, 2, 4, 7, 10, 13, 16]]
    add_table(doc, ["交会角", "后验直径"], [(f"{x['crossing_angle_deg']:.0f}°", f"{x['diameter_m']:.3f} m") for x in chosen], widths=[7.5, 7.6], caption="表6  代表性交会角下的后验直径")
    add_body(doc, "10°时直径为354.620 m，90°时降至42.626 m，下降约88.0%。因此“向首条射线侧向移动”具有明确的几何依据；但90°并非独立于覆盖条件的硬目标，实际最优点仍由式（13）的稳健排序决定。")

    add_heading(doc, "7.2  离散精度与计算复杂度", 2)
    add_body(doc, "圆弧正多边形边数N越大，弦高误差按O(N⁻²)下降。对固定N，m次观测的半平面裁剪与当前顶点数近似线性相关；得到n个凸多边形顶点后，旋转卡壳为O(n)，最小包围圆为O(n)期望时间。第二问若有M个候选点、K个源假设和L个误差样本，总体约为O(MKL·n)。由于各候选点独立，最外层可以直接并行。")
    add_body(doc, "实际计算应对N、网格间距和误差离散数做收敛检查：逐次加倍N、减半候选网格间距，并比较最优点、最坏直径和推荐区域是否稳定。该检查服务于数值可信度，不改变题面模型。")

    add_heading(doc, "7.3  模型边界", 2)
    add_bullet(doc, "本文把机器人坐标视为精确；若存在位置误差，应把Sᵢ的不确定性并入集合传播，否则定位区域会偏小。")
    add_bullet(doc, "字典序目标适合强调任务安全的清除场景；若赛事后续给定明确时间代价，可在保证漏测率阈值后再做多目标Pareto比较。")
    add_bullet(doc, "候选区基于全向传播和圆形覆盖；山体遮挡、多径和障碍物应作为可达域或可见域掩膜加入，而不能通过调大角度误差代替。")

    add_heading(doc, "8  模型评价", 1)
    add_heading(doc, "8.1  优点", 2)
    add_bullet(doc, "严格尊重±1°硬误差和定点误差不变的题设，不依赖缺乏证据的概率分布。")
    add_bullet(doc, "定位区域始终保持凸性，直径和最小包围圆均有成熟的线性或线性期望算法，结论可证明、计算可复现。")
    add_bullet(doc, "把“是否覆盖”从直觉判断转化为R*=D/2的必要充分条件，并给出反例，完整回答了问题一的追问。")
    add_bullet(doc, "第二问同时处理可靠接收、最坏定位精度、平均表现和路程，候选区域可直接交给机器人规划模块。")
    add_heading(doc, "8.2  不足与改进", 2)
    add_body(doc, "主要不足是第二问仍需离散Ω₁和候选空间，结果受网格分辨率影响；可用自适应四叉树、贝叶斯优化或连续非光滑优化进一步减少计算量。当前模型也未显式考虑地形传播和移动障碍，后续可把可通行区域、视距和实测信号强度作为额外集合约束。")

    add_heading(doc, "9  结论", 1, page_break=True)
    add_body(doc, "问题一的定位区域是目标圆域、各示向误差扇形与最大接收圆盘的凸交集。圆弧离散并裁剪后，使用旋转卡壳即可得到区域直径；直径为D的圆能覆盖区域的充要条件是其最小包围圆半径R*=D/2，一般凸多边形并不满足该条件。")
    add_body(doc, "问题二应把第二测点选择建模为覆盖约束下的稳健几何优化。推荐点首先应在1000 m最小有效半径下尽量覆盖首次可行域，其次形成较大的交会角，并使所有源假设和误差端点下的最大后验直径最小。输出最优点之外，还应输出最坏直径位于前20%的零漏测候选区域，以便机器人结合真实障碍和路径代价执行。")

    add_heading(doc, "参考文献", 1, page_break=True)
    refs = [
        "[1] 全国大学生数学建模竞赛组委会. 2026年高教社杯全国大学生数学建模竞赛B题：无线电干扰源的快速自动定位与清除. 2026.",
        "[2] Kaya C Y. Observer path planning for maximum information. Optimization, 2022, 71(4): 1097–1116. DOI: 10.1080/02331934.2021.2011868.",
        "[3] Zou Y, Gao B, Tang X, Yu L. Target localization and sensor movement trajectory planning with bearing-only measurements in three dimensional space. Applied Sciences, 2022, 12(13): 6739. DOI: 10.3390/app12136739.",
        "[4] Zhou R, et al. Optimal 3D angle of arrival sensor placement with Gaussian priors. Entropy, 2021, 23(11): 1379. DOI: 10.3390/e23111379.",
        "[5] Toussaint G T. Solving geometric problems with the rotating calipers. Proceedings of IEEE MELECON, 1983: A10.02/1–4.",
        "[6] Welzl E. Smallest enclosing disks (balls and ellipsoids). In: New Results and New Trends in Computer Science. LNCS 555. Springer, 1991: 359–370. DOI: 10.1007/BFb0038202.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(ref)
        set_run_font(r, size=9.5)
        set_keep_lines(p)

    add_heading(doc, "附录A  核心算法伪代码", 1, page_break=True)
    add_heading(doc, "A.1  问题一直径与覆盖判定", 2)
    algo1 = [
        ("输入", "检测点Sᵢ、读数θ̂ᵢ、误差ε、圆弧离散精度N"),
        ("初始化", "P←目标圆Ω₀的正N边形"),
        ("循环1", "对每个Sᵢ，以θ̂ᵢ±ε构造两个半平面，P←Clip(P,Hᵢ⁻)，P←Clip(P,Hᵢ⁺)"),
        ("循环2", "P←P∩B(Sᵢ,1500)；若P为空则报告观测不一致"),
        ("直径", "D,(a,b)←RotatingCalipers(P)"),
        ("包围圆", "(c,R*)←MinimumEnclosingCircle(P)"),
        ("输出", "P、D、端点(a,b)，以及R*≤D/2+τ的覆盖判定"),
    ]
    add_table(doc, ["行", "伪代码"], algo1, widths=[2.0, 13.1], caption="表A1  交会定位与覆盖判定")
    add_heading(doc, "A.2  第二测点稳健选址", 2)
    algo2 = [
        ("输入", "首次可行域Ω₁、检测点S₁、候选网格C、源假设G、误差集合E"),
        ("对每个s∈C", "统计‖g−s‖>1000的假设并计算pₘᵢₛₛ(s)"),
        ("可示向假设", "对每个g∈G和δ∈E构造Ω₂(s,g,δ)，调用问题一算法求D(Ω₂)"),
        ("聚合", "求Dmax(s)、Dmean(s)与‖s−S₁‖"),
        ("排序", "按(pₘᵢₛₛ,Dmax,Dmean,路程)字典序升序"),
        ("区域", "在pₘᵢₛₛ最小的点中保留Dmax前20%，形成推荐区域Cρ"),
        ("输出", "最优点、候选区域、每点四项指标与未覆盖子区"),
    ]
    add_table(doc, ["行", "伪代码"], algo2, widths=[2.0, 13.1], caption="表A2  第二测点稳健选址")

    # Document properties and save.
    doc.core_properties.title = "有界示向误差下的无线电干扰源交会定位与稳健观测点设计"
    doc.core_properties.subject = "2026 CUMCM B题问题一与问题二"
    doc.core_properties.author = "参赛队论文草稿"
    doc.core_properties.keywords = "数学建模, 示向定位, 凸几何, 稳健优化"
    doc.core_properties.comments = "由题面约束、可复现几何算法与公开文献共同形成的阶段稿。"
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
